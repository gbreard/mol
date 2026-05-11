#!/usr/bin/env python3
"""
Genera exports/comparacion_esco_vs_mol_diego.xlsx — comparación lado a lado
ESCO catálogo vs lo que MOL extrajo, para las 4 ocupaciones que reportó Diego.

1 hoja consolidada con todas las skills, marcando si vienen de:
  - ESCO_only:  catalogada por ESCO pero NO extraída por MOL
  - MOL_only:   extraída por MOL pero NO en catálogo ESCO de la ocupación
  - both:       presente en ambos
"""

import sqlite3
import json
import statistics
from pathlib import Path
from collections import defaultdict

import openpyxl

PROJECT = Path("/mnt/d/OEDE/Webscrapping")
DB_PATH = f"file:{PROJECT}/database/bumeran_scraping.db?mode=ro"

ESCO_OCC_FULL_JSON     = PROJECT / "database/embeddings/esco_occupations_full.json"
ESCO_OCC_SKILLS_JSON   = PROJECT / "database/embeddings/esco_occupation_skills.json"
ESCO_SKILLS_META_JSON  = PROJECT / "database/embeddings/esco_skills_metadata_full.json"
ESCO_RDF_EXTRA_JSON    = PROJECT / "exports/esco_skills_rdf_extract.json"

OUTPUT_XLSX = PROJECT / "exports/comparacion_esco_vs_mol_diego.xlsx"

# 4 ocupaciones reportadas por Diego (label en BD, ISCO 4d)
TARGETS = [
    ('representante comercial',                              '3322'),
    ('vendedor especializado/vendedora especializada',       '5223'),
    ('empleado de oficina/empleada de oficina',              '4110'),
    ('desarrollador de software/desarrolladora de software', '2512'),
]

REUSE_LABEL_ES = {
    "transversal": "Transversal",
    "cross-sector": "Inter-sectorial",
    "sector-specific": "Específica del sector",
    "occupation-specific": "Específica de la ocupación",
}

ID_CAP = 200


def fmt_ids(ids_list, total_count):
    if not ids_list:
        return ""
    head = ";".join(str(x) for x in ids_list[:ID_CAP])
    if total_count > ID_CAP:
        return f"{head};...+{total_count - ID_CAP} más"
    return head


def main():
    print("Cargando catálogos ESCO...")
    with open(ESCO_OCC_FULL_JSON) as f:
        esco_occ = json.load(f)
    occs_by_uri = {o['uri']: o for o in esco_occ['occupations']}

    with open(ESCO_SKILLS_META_JSON) as f:
        skills_meta = json.load(f)
    skill_by_uri = {s['uri']: s for s in skills_meta}

    with open(ESCO_OCC_SKILLS_JSON) as f:
        occ_skills_data = json.load(f)
    occ_skills_map = occ_skills_data['occupation_skills']

    with open(ESCO_RDF_EXTRA_JSON) as f:
        rdf_extra = json.load(f)
    skills_reuse = rdf_extra['skills_reuse']
    skills_green = set(rdf_extra['skills_green'])

    print("Conectando BD MOL (read-only)...")
    con = sqlite3.connect(DB_PATH, uri=True)
    cur = con.cursor()

    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Comparacion ESCO vs MOL")
    ws.append([
        "Ocupación reportada", "ISCO 4d",
        "ESCO URI ocupación", "ESCO code ocupación",
        "Skill label", "Skill URI",
        "L1 código", "L2 código",
        "Reuse level", "Green",
        "Estado",                       # ESCO_only / MOL_only / both
        "Relación ESCO con la ocup.",   # Esencial / Optativa / —
        "N apariciones MOL",
        "Cobertura % MOL",
        "Score promedio MOL",
        "Posición ranking MOL (en la ocup.)",
        f"id_ofertas (cap {ID_CAP})",
    ])

    total_rows = 0

    for label_target, isco_target in TARGETS:
        print(f"\n--- {label_target} (ISCO {isco_target}) ---")

        # 1) Encontrar ofertas de esa ocupación
        offer_uris = set()
        n_offers = 0
        for row in cur.execute("""
            SELECT id_oferta, esco_occupation_uri
            FROM ofertas_esco_matching
            WHERE esco_occupation_label = ? AND isco_code = ?
              AND esco_occupation_uri IS NOT NULL
        """, (label_target, isco_target)):
            offer_uris.add(row[0])
            n_offers += 1

        # Determinar la ESCO URI más frecuente para esta ocupación reportada
        # (puede haber varias por mismatches)
        uri_counts = defaultdict(int)
        for row in cur.execute("""
            SELECT esco_occupation_uri, COUNT(*) n
            FROM ofertas_esco_matching
            WHERE esco_occupation_label = ? AND isco_code = ?
              AND esco_occupation_uri IS NOT NULL
            GROUP BY esco_occupation_uri
            ORDER BY n DESC
        """, (label_target, isco_target)):
            uri_counts[row[0]] += row[1]
        if not uri_counts:
            print(f"  Sin ofertas matcheadas, salteando")
            continue
        primary_uri = max(uri_counts, key=uri_counts.get)
        primary_meta = occs_by_uri.get(primary_uri, {})
        print(f"  ESCO URI principal: {primary_meta.get('esco_code')}  ({len(uri_counts)} URIs distintas)")
        print(f"  Ofertas: {n_offers}")

        # 2) Catálogo ESCO essential+optional para la ocupación principal
        esco_skills = {}  # skill_uri -> 'essential' | 'optional'
        cat = occ_skills_map.get(primary_uri, {})
        for s in cat.get('essential', []):
            esco_skills[s['skill_uri']] = 'essential'
        for s in cat.get('optional', []):
            esco_skills[s['skill_uri']] = 'optional'
        print(f"  Skills en catálogo ESCO: {len(esco_skills)}")

        # 3) Skills extraídas por MOL para esas ofertas
        mol_skills = defaultdict(lambda: {'count': 0, 'scores': [], 'label': '', 'ids': []})
        if offer_uris:
            placeholders = ','.join('?' * len(offer_uris))
            for row in cur.execute(f"""
                SELECT id_oferta, esco_skill_uri, esco_skill_label, match_score
                FROM ofertas_esco_skills_detalle
                WHERE id_oferta IN ({placeholders})
                  AND esco_skill_uri IS NOT NULL AND esco_skill_uri != ''
            """, list(offer_uris)):
                id_oferta, skill_uri, skill_label, score = row
                if not skill_label:
                    skill_label = skill_by_uri.get(skill_uri, {}).get('label', '')
                if not skill_label:
                    continue
                d = mol_skills[skill_uri]
                d['count'] += 1
                d['label'] = skill_label
                if score is not None:
                    d['scores'].append(score)
                if len(d['ids']) < ID_CAP:
                    d['ids'].append(id_oferta)
        print(f"  Skills extraídas MOL: {len(mol_skills)}")

        # 4) Ranking MOL por frecuencia
        sorted_mol = sorted(mol_skills.items(), key=lambda x: -x[1]['count'])
        mol_rank = {su: i + 1 for i, (su, _) in enumerate(sorted_mol)}

        # 5) Construir filas: union de ambos sets
        all_uris = set(esco_skills) | set(mol_skills)

        for skill_uri in all_uris:
            in_esco = skill_uri in esco_skills
            in_mol = skill_uri in mol_skills

            if in_esco and in_mol:
                estado = "both"
            elif in_esco:
                estado = "ESCO_only"
            else:
                estado = "MOL_only"

            sm = skill_by_uri.get(skill_uri, {})
            label = sm.get('label', '') or mol_skills.get(skill_uri, {}).get('label', '')

            relation = esco_skills.get(skill_uri)
            relation_es = {'essential': 'Esencial', 'optional': 'Optativa'}.get(relation, '—')

            mol_data = mol_skills.get(skill_uri, {'count': 0, 'scores': [], 'label': '', 'ids': []})
            mol_count = mol_data['count']
            mol_cov = round(mol_count / n_offers * 100, 2) if n_offers and mol_count else 0
            mol_score = round(statistics.mean(mol_data['scores']), 3) if mol_data['scores'] else 0
            mol_rank_val = mol_rank.get(skill_uri, '')
            ids_str = fmt_ids(mol_data['ids'], mol_count)

            reuse_raw = skills_reuse.get(skill_uri, '')
            reuse_label = REUSE_LABEL_ES.get(reuse_raw, reuse_raw)
            green = "Sí" if skill_uri in skills_green else "No"

            ws.append([
                label_target, isco_target,
                primary_uri, primary_meta.get('esco_code', ''),
                label, skill_uri,
                sm.get('L1', ''), sm.get('L2', ''),
                reuse_label, green,
                estado, relation_es,
                mol_count, mol_cov, mol_score, mol_rank_val,
                ids_str,
            ])
            total_rows += 1

    con.close()

    OUTPUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_XLSX)
    size_kb = OUTPUT_XLSX.stat().st_size / 1024
    print(f"\nGuardado: {OUTPUT_XLSX} ({size_kb:.1f} KB)")
    print(f"Total filas: {total_rows}")


if __name__ == "__main__":
    main()
