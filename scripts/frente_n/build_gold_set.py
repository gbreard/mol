# -*- coding: utf-8 -*-
"""
[FRENTE N — N4] Construye el gold set de tareas desde las marcas de Cyn (28 casos).

gold_tareas por caso = (extraídas correctas, no objetadas) + FALTA (redacción de
Cyn) − SOBRA. Granularidad = la de Cyn (unidad funcional). Se encodean a mano
desde su validación (fuente: validacion_tareas_respuestas_2026-08-26.xlsx) porque
FALTA/SOBRA son prosa libre; los campos estructurales se leen del Excel.

Salida: metrics/gold_set_tareas.json (protegido; repo privado — patrón del gold
set existente database/gold_set_manual_v2.json).
"""
import json
import openpyxl
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / 'exports/cyn_backlog/validacion_tareas_respuestas_2026-08-26.xlsx'
OUT = ROOT / 'metrics/gold_set_tareas.json'

# gold_tareas por id (redacción de Cyn, granularidad de unidad funcional).
# [] = salida vacía válida (los OK sin tareas y los que ella dejó en cero).
GOLD = {
    '1118165658': [  # C1 FALLA
        "analizar y evaluar riesgo crediticio para el otorgamiento de líneas de crédito a clientes",
        "realizar seguimiento de pedidos",
        "coordinar con el área comercial la correcta entrega y cobranza de los pedidos",
        "controlar y realizar seguimiento de cuentas corrientes comerciales",
        "gestionar integralmente la facturación",
        "realizar seguimiento de cobranzas",
        "controlar vencimientos",
        "elaborar reportes de crédito, facturación y morosidad",
        "trabajar coordinadamente con las áreas Comercial, Administración y Finanzas",
    ],
    '5790086984': [  # C2 FALLA + sobra huérfano
        "diseñar, planificar y ejecutar estrategias integrales de testing",
        "crear y mantener casos de prueba, planes de testing y documentación técnica detallada",
        "identificar, reportar y dar seguimiento a los bugs hasta su resolución",
        "asegurar que el producto cumpla con los requisitos funcionales y no funcionales (performance, seguridad y usabilidad)",
        "participar de las ceremonias ágiles",
        "articular activamente con desarrolladores, BAs y otros stakeholders",
        "aportar análisis crítico sobre los requerimientos",
        "actuar como nexo entre los equipos técnicos y los clientes",
    ],
    '1118392033': [  # C3 FALLA
        "realizar la carga, descarga y reparto de mercadería a clientes",
        "brindar una atención cordial a los clientes",
        "cuidar el vehículo y la mercadería asignada",
    ],
    '2176458': [  # C4 FALLA
        "mecanizar en máquinas CNC con programación a pie de máquina mediante sistema FAGOR",
        "usar y mantener herramientas de medición",
        "ejecutar las órdenes de trabajo",
        "controlar la entrada y salida al taller de materiales y equipos de trabajo",
        "llevar registro y control de los trabajos realizados y/o a realizar en el taller",
        "mantener en orden el equipo y el sitio de trabajo, reportando cualquier anomalía",
        "cumplir las normas y procedimientos en materia de seguridad integral",
    ],
    '1117972807': [  # C5 FALLA
        "colaborar con la liquidación de sueldos",
        "tareas de contabilidad general",
        "tareas generales de oficina",
    ],
    '1118316198': [  # C6 FALLA
        "planificar, ejecutar y optimizar la estrategia de performance digital",
        "impulsar la adquisición de usuarios",
        "monitorear y analizar indicadores clave como CPA, ROAS, conversión, retención y calidad de tráfico",
        "identificar oportunidades de mejora a partir del análisis de datos y comportamiento de usuarios",
    ],
    '2176705': [  # C7 FALLA
        "diseñar, desarrollar y entregar soluciones de software confiables y de alta calidad",
        "contribuir al desarrollo de código productivo y participar activamente en revisiones de código dentro del equipo",
        "resolver problemas técnicos complejos relacionados con el diseño, el desarrollo y la estabilidad de las aplicaciones",
        "identificar y ejecutar oportunidades de automatización para reducir incidentes recurrentes y mejorar la estabilidad operativa",
        "participar en evaluaciones técnicas con proveedores externos, startups y equipos internos cuando sea requerido",
        "promover el uso de nuevas tecnologías, herramientas y buenas prácticas dentro de la comunidad de ingeniería",
        "contribuir a una cultura de diversidad, equidad e inclusión dentro del equipo y la organización",
    ],
    '6477274334': [],  # C8 FALLA (sobra 2) → cero tareas
    '1118058753': [],  # C9 FALLA (sobra 2) → cero tareas
    '6265862938': [],  # C10 FALLA (sobra 1) → cero tareas
    '6171358600': [],  # C11 FALLA (sobra 2) → cero tareas
    '2185030': [  # C12 FALLA
        "captar nuevos asociados",
        "comercializar planes de cobertura médica",
        "realizar prospección activa presencial y/o digital de nuevos clientes individuales y corporativos, gestionando el proceso completo de venta: presupuestación por segmento, cierre y documentación para el alta",
        "recibir y gestionar leads provenientes de distintos canales",
        "registrar y realizar seguimiento de la gestión comercial en CRM",
        "analizar la competencia y reportar acciones del mercado",
    ],
    '5273316732': [  # C13 FALLA
        "ofrecer productos bancarios",
        "comunicarse de manera clara, efectiva y persuasiva para generar ventas",
    ],
    '5917821361': [  # C14 FALLA
        "brindar atención cordial, eficiente y resolutiva a los clientes en sala y vía pública",
        "promocionar los servicios, eventos y beneficios del Bingo",
        "garantizar una experiencia satisfactoria y segura bajo los protocolos de servicio",
        "colaborar en acciones de marketing y fidelización de clientes",
    ],
    '5650193558': [],  # C15 FALLA (sobra 1) → cero tareas
    '2176546': [  # C16 OK (5 tareas, se conservan tal cual)
        "realizar extracciones de sangre para la obtención de plasma rico en plaquetas (PRP)",
        "administrar sueros vitamínicos según protocolos establecidos",
        "brindar atención y orientación a los pacientes durante los procedimientos",
        "mantener un ambiente de trabajo seguro y estéril",
        "colaborar con el equipo para asegurar la calidad del servicio",
    ],
    '2169530': [  # C17 FALLA
        "realizar el mantenimiento preventivo y correctivo de una flota vehicular, incluyendo máquinas viales, camiones y utilitarios",
        "reparar maquinarias pesadas, camiones y utilitarios",
        "diagnosticar y resolver fallas mecánicas en sistemas hidráulicos y neumáticos",
        "realizar reparaciones generales y trabajos de soldadura",
    ],
    '2168972': [  # C18 FALLA (sobra 1)
        "asistir a audiencias",
        "realizar el seguimiento de la cartera de clientes",
    ],
    '6097371222': [],  # C19 OK → cero tareas
    '2172514': [],     # C20 OK → cero tareas
    '5909299836': [],  # C21 OK → cero tareas
    '5160083235': [    # C22 OK (1 tarea)
        "reponer los productos en góndolas",
    ],
    '6483882943': [    # C23 OK (1 tarea)
        "realizar la carga y descarga de mercadería",
    ],
    '2187269': [  # C24 FALLA (gemelo de C12)
        "captar nuevos asociados",
        "comercializar planes de cobertura médica",
        "realizar prospección activa presencial y/o digital de nuevos clientes individuales y corporativos, gestionando el proceso completo de venta: presupuestación por segmento, cierre y documentación para el alta",
        "recibir y gestionar leads provenientes de distintos canales",
        "registrar y realizar seguimiento de la gestión comercial en CRM",
        "analizar la competencia",
        "reportar acciones del mercado",
    ],
    '1118076732': [],  # C25 FALLA (sobra 2) → cero tareas
    '2181056': [],     # C26 FALLA (sobra 1) → cero tareas
    '1118300373': [  # C27 FALLA
        "aplicar tratamientos de medicina regenerativa y estética ginecológica",
        "participar en la formación y capacitación en protocolos exclusivos",
        "brindar atención profesional, cálida y empática a las pacientes",
        "contribuir a un ambiente de excelencia e innovación en medicina estética",
    ],
    '2176620': [  # C28 FALLA
        "desarrollar, configurar y mantener pantallas HMI en Fast Tools (Yokogawa) para sistemas SCADA de producción y facilities",
        "programar y configurar gráficos de proceso, tendencias y alarmas en DeltaV Operate para plantas de tratamiento y baterías de producción",
        "diseñar arquitecturas de visualización SCADA integradas, asegurando consistencia entre sistemas Fast Tools y DeltaV",
        "implementar estándares de HMI según ISA-101 y guías de alto rendimiento (High Performance HMI)",
        "coordinar con Operaciones para definir requerimientos de visualización y optimizar interfaces de operador",
        "elaborar documentación técnica de pantallas SCADA (especificaciones funcionales, narrativas de operación, manuales de usuario)",
    ],
}

# Sobras TIPIFICADAS que Cyn marcó explícitamente (para el chequeo "cero sobras").
SOBRAS_TIPIFICADAS = {
    '5790086984': ["seguridad y usabilidad"],                 # fragmento huérfano
    '6477274334': ["atención al cliente", "disponibilidad para trabajar fines de semana"],
    '1118058753': ["realizar selección y administración de personal", "disponibilidad para trabajar presencial"],
    '6265862938': ["atención al público"],
    '6171358600': ["realizar dibujos y planos", "manejar autocad"],
    '5650193558': ["manejo de maquinas de limpieza"],
    '2168972': ["trabajar en conjunto con un equipo comprometido y en constante crecimiento"],
    '1118076732': ["asesoramiento financiero", "normativa y procesos de prevención de lavado de activos"],
    '2181056': ["manejo de protocolo"],
}


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb['28 casos']
    casos = []
    for r in range(2, ws.max_row + 1):
        g = lambda c: ws[f'{c}{r}'].value
        oid = str(g('B')).strip() if g('B') is not None else None
        if not oid:
            continue
        casos.append({
            'n': g('A'),
            'id_oferta': oid,
            'portal': g('C'),
            'titulo': g('D'),
            'cuerpo': g('E'),
            'sistema_v11_extrajo': g('F'),
            'veredicto_cyn': (g('I') or '').strip(),
            'gold_tareas': GOLD.get(oid, []),
            'sobras_tipificadas': SOBRAS_TIPIFICADAS.get(oid, []),
        })
    faltan = [c['id_oferta'] for c in casos if c['id_oferta'] not in GOLD]
    assert not faltan, f'IDs sin gold encodeado: {faltan}'
    doc = {
        'version': '1.0',
        'fuente': 'validacion_tareas_respuestas_2026-08-26.xlsx (28 casos, Cyn)',
        'nota': 'gold_tareas a granularidad de unidad funcional (método Cyn). [] = vacío válido.',
        'n_casos': len(casos),
        'n_ok': sum(1 for c in casos if c['veredicto_cyn'].upper().startswith('OK')),
        'casos': casos,
    }
    OUT.write_text(json.dumps(doc, ensure_ascii=False, indent=2))
    tot_gold = sum(len(c['gold_tareas']) for c in casos)
    vac = sum(1 for c in casos if not c['gold_tareas'])
    print(f'gold set -> {OUT}')
    print(f'  casos: {len(casos)}  tareas-oro totales: {tot_gold}  casos vacío-válido: {vac}')
    print(f'  OK (6 intocables esperados): {doc["n_ok"]}')


if __name__ == '__main__':
    main()
